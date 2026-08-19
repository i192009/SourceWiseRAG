from pathlib import Path
from typing import Any, Dict, List

import pandas as pd
from PyPDF2 import PdfReader
from docx import Document


class ParserService:
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        suffix = Path(file_path).suffix.lower()

        if suffix == ".pdf":
            return self._parse_pdf(file_path)
        if suffix == ".docx":
            return self._parse_docx(file_path)
        if suffix in {".xlsx", ".xls"}:
            return self._parse_excel(file_path)
        if suffix in {".png", ".jpg", ".jpeg"}:
            return self._parse_image(file_path)
        if suffix in {".txt", ".md"}:
            return self._parse_text(file_path)

        raise ValueError(f"Unsupported file type: {suffix}")

    def _parse_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        reader = PdfReader(file_path)
        units: List[Dict[str, Any]] = []

        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""

            units.append({
                "unit_type": "pdf_page",
                "unit_index": i,
                "page_number": i + 1,
                "sheet_name": None,
                "heading_path": [],
                "raw_text": text,
                "needs_ocr": len(text.strip()) == 0,
                "ocr_confidence": None,
                "source_meta": {
                    "page_index": i,
                    "file_path": file_path,
                },
            })

        return units

    def _parse_docx(self, file_path: str) -> List[Dict[str, Any]]:
        doc = Document(file_path)
        units: List[Dict[str, Any]] = []

        heading_stack: List[str] = []
        unit_index = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ""

            if style_name.startswith("Heading"):
                level = self._extract_heading_level(style_name)
                if level is not None:
                    heading_stack = heading_stack[: level - 1]
                    heading_stack.append(text)

                    units.append({
                        "unit_type": "docx_heading",
                        "unit_index": unit_index,
                        "page_number": None,
                        "sheet_name": None,
                        "heading_path": heading_stack.copy(),
                        "raw_text": text,
                        "needs_ocr": False,
                        "ocr_confidence": None,
                        "source_meta": {
                            "style_name": style_name,
                            "heading_level": level,
                        },
                    })
                    unit_index += 1
                    continue

            units.append({
                "unit_type": "docx_paragraph",
                "unit_index": unit_index,
                "page_number": None,
                "sheet_name": None,
                "heading_path": heading_stack.copy(),
                "raw_text": text,
                "needs_ocr": False,
                "ocr_confidence": None,
                "source_meta": {
                    "style_name": style_name,
                },
            })
            unit_index += 1

        for table_idx, table in enumerate(doc.tables):
            rows = []
            for row_idx, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    row_text = " | ".join(
                        f"col_{i + 1}: {value}" for i, value in enumerate(cells) if value
                    )
                    rows.append(f"row_{row_idx + 1}: {row_text}")

            if rows:
                units.append({
                    "unit_type": "docx_table",
                    "unit_index": unit_index,
                    "page_number": None,
                    "sheet_name": f"table_{table_idx + 1}",
                    "heading_path": heading_stack.copy(),
                    "raw_text": "\n".join(rows),
                    "needs_ocr": False,
                    "ocr_confidence": None,
                    "source_meta": {
                        "table_index": table_idx,
                    },
                })
                unit_index += 1

        return units

    def _parse_excel(self, file_path: str) -> List[Dict[str, Any]]:
        sheets = pd.read_excel(file_path, sheet_name=None)
        units: List[Dict[str, Any]] = []

        for sheet_idx, (sheet_name, df) in enumerate(sheets.items()):
            df = df.fillna("")
            row_texts = []

            for row_idx, row in df.iterrows():
                cells = []
                for col_name, value in row.items():
                    value = str(value).strip()
                    if value:
                        cells.append(f"{col_name}: {value}")

                if cells:
                    row_texts.append(f"row_{row_idx + 1}: " + " | ".join(cells))

            units.append({
                "unit_type": "excel_sheet",
                "unit_index": sheet_idx,
                "page_number": None,
                "sheet_name": str(sheet_name),
                "heading_path": [],
                "raw_text": "\n".join(row_texts),
                "needs_ocr": False,
                "ocr_confidence": None,
                "source_meta": {
                    "sheet_index": sheet_idx,
                    "row_count": len(row_texts),
                    "column_names": [str(c) for c in df.columns],
                },
            })

        return units

    def _parse_image(self, file_path: str) -> List[Dict[str, Any]]:
        return [{
            "unit_type": "image_ocr",
            "unit_index": 0,
            "page_number": 1,
            "sheet_name": None,
            "heading_path": [],
            "raw_text": "",
            "needs_ocr": True,
            "ocr_confidence": None,
            "source_meta": {
                "image_index": 0,
                "file_path": file_path,
            },
        }]

    def _parse_text(self, file_path: str) -> List[Dict[str, Any]]:
        text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
        return [{
            "unit_type": "text_file",
            "unit_index": 0,
            "page_number": None,
            "sheet_name": None,
            "heading_path": [],
            "raw_text": text,
            "needs_ocr": False,
            "ocr_confidence": None,
            "source_meta": {
                "file_path": file_path,
            },
        }]

    def _extract_heading_level(self, style_name: str):
        parts = style_name.split()
        if len(parts) == 2 and parts[0] == "Heading" and parts[1].isdigit():
            return int(parts[1])
        return None
